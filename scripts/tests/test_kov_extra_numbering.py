from io import BytesIO
from pathlib import Path
import sys
import unittest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from scrape_secondary_curricula import SourceDocument,parse_kov_docx
from secondary_record_schema import normalize_curriculum_record

class KovExtraNumberingTests(unittest.TestCase):
    def document(self, numbered=True):
        d=Document()
        for name in ['Doel','Doel: Extra','MD + SMD + BK']:d.styles.add_style(name,1)
        if numbered:
            abstract=OxmlElement('w:abstractNum');abstract.set(qn('w:abstractNumId'),'700')
            level=OxmlElement('w:lvl');level.set(qn('w:ilvl'),'0')
            for tag,value in [('start','2'),('numFmt','decimal'),('lvlText','LPD %1 +')]:
                e=OxmlElement('w:'+tag);e.set(qn('w:val'),value);level.append(e)
            abstract.append(level);d.part.numbering_part.element.append(abstract)
            num=OxmlElement('w:num');num.set(qn('w:numId'),'700');ref=OxmlElement('w:abstractNumId');ref.set(qn('w:val'),'700');num.append(ref);d.part.numbering_part.element.append(num)
            props=d.styles['Doel: Extra'].element.get_or_add_pPr();n=OxmlElement('w:numPr');ref=OxmlElement('w:numId');ref.set(qn('w:val'),'700');n.append(ref);props.append(n)
        d.add_heading('Communicatie',level=2)
        d.add_paragraph('MD 02. 08 Synthetisch minimumdoel. (LPD 3)',style='MD + SMD + BK')
        d.add_paragraph('MD 09.01 Ander regulier minimumdoel. (LPD 2)',style='MD + SMD + BK')
        d.add_paragraph('De leerlingen lezen een tekst.',style='Doel')
        d.add_paragraph('De leerlingen formuleren een eigen mening.',style='Doel: Extra')
        d.add_heading('Literatuur',level=2)
        d.add_paragraph('De leerlingen vertellen over een boek.',style='Doel')
        return d

    def parse(self,d):
        b=BytesIO();d.save(b)
        return parse_kov_docx(b.getvalue(),SourceDocument('KOV','Synthetic','https://example.test/plan.docx','https://example.test/I-Test-a/leerplan','1ste graad',stream='A-stroom',discipline='Nederlands'))

    def test_extra_keeps_source_label_and_later_minimum_links(self):
        rs=self.parse(self.document());self.assertEqual([r.code for r in rs],['I-Test-a LPD 1','I-Test-a LPD 2 +','I-Test-a LPD 3'])
        self.assertEqual(rs[1].subdomein,'Communicatie');self.assertEqual(rs[1].minimumdoel_codes,[])
        self.assertEqual(rs[2].subdomein,'Literatuur');self.assertEqual(rs[2].minimumdoel_codes,['02.08'])

    def test_unproven_extra_numbering_is_rejected(self):
        with self.assertRaisesRegex(ValueError,'nummering'):self.parse(self.document(False))

    def test_export_preserves_context_and_minimum_links(self):
        rs=self.parse(self.document());r=normalize_curriculum_record(rs[-1])
        self.assertEqual(r.get('subdomein'),'Literatuur');self.assertEqual(r.get('minimumdoel_codes'),['02.08']);self.assertEqual(r.get('bron_titel'),'Synthetic')

    def test_conflicting_explicit_restart_is_rejected(self):
        d=self.document()
        num=OxmlElement('w:num');num.set(qn('w:numId'),'701')
        ref=OxmlElement('w:abstractNumId');ref.set(qn('w:val'),'700');num.append(ref)
        override=OxmlElement('w:lvlOverride');override.set(qn('w:ilvl'),'0')
        start=OxmlElement('w:startOverride');start.set(qn('w:val'),'9');override.append(start);num.append(override)
        d.part.numbering_part.element.append(num)
        props=d.paragraphs[-1]._p.get_or_add_pPr();n=OxmlElement('w:numPr')
        ref=OxmlElement('w:numId');ref.set(qn('w:val'),'701');n.append(ref);props.append(n)
        with self.assertRaisesRegex(ValueError,'herstart'):self.parse(d)

    def regular_document(self, instances, base_start="1"):
        d = Document()
        d.styles.add_style('Doel', 1)
        abstract = OxmlElement('w:abstractNum')
        abstract.set(qn('w:abstractNumId'), '800')
        level = OxmlElement('w:lvl'); level.set(qn('w:ilvl'), '0')
        for tag, value in [('start', base_start), ('numFmt', 'decimal'), ('lvlText', 'LPD %1')]:
            if value is not None:
                element = OxmlElement('w:' + tag); element.set(qn('w:val'), value); level.append(element)
        abstract.append(level); d.part.numbering_part.element.append(abstract)
        for num_id, start in dict(instances).items():
            num = OxmlElement('w:num'); num.set(qn('w:numId'), str(num_id))
            ref = OxmlElement('w:abstractNumId'); ref.set(qn('w:val'), '800'); num.append(ref)
            if start is not None:
                override = OxmlElement('w:lvlOverride'); override.set(qn('w:ilvl'), '0')
                value = OxmlElement('w:startOverride'); value.set(qn('w:val'), str(start)); override.append(value)
                num.append(override)
            d.part.numbering_part.element.append(num)
        for num_id, _ in instances:
            paragraph = d.add_paragraph('De leerlingen lezen.', style='Doel')
            props = paragraph._p.get_or_add_pPr(); numbering = OxmlElement('w:numPr')
            level = OxmlElement('w:ilvl'); level.set(qn('w:val'), '0'); numbering.append(level)
            ref = OxmlElement('w:numId'); ref.set(qn('w:val'), str(num_id)); numbering.append(ref)
            props.append(numbering)
        return d

    def test_same_explicit_numbering_instance_continues(self):
        records = self.parse(self.regular_document([(801, 1), (801, 1)]))
        self.assertEqual([r.code for r in records], ['I-Test-a LPD 1', 'I-Test-a LPD 2'])

    def test_separate_instances_with_matching_restarts_continue(self):
        records = self.parse(self.regular_document([(801, 1), (801, 1), (802, 3), (802, 3)]))
        self.assertEqual([r.code for r in records], [f'I-Test-a LPD {n}' for n in range(1, 5)])

    def test_separate_instance_with_conflicting_restart_is_rejected(self):
        with self.assertRaisesRegex(ValueError, 'herstart'):
            self.parse(self.regular_document([(801, 1), (802, 1)]))

    def test_return_to_interrupted_instance_is_rejected(self):
        with self.assertRaisesRegex(ValueError, 'lijstvoortzetting'):
            self.parse(self.regular_document([(801, 1), (802, 2), (801, 1)]))

    def test_numbering_state_is_local_to_document(self):
        for _ in range(2):
            self.assertEqual(len(self.parse(self.regular_document([(801, 1)]))), 1)

    def test_unconfirmed_regular_numbering_is_rejected(self):
        for tag, value in [('numFmt', 'upperRoman'), ('lvlText', 'Other %1')]:
            with self.subTest(tag=tag):
                document = self.regular_document([(801, 1)])
                abstract = next(a for a in document.part.numbering_part.element.findall(qn('w:abstractNum')) if a.get(qn('w:abstractNumId')) == '800')
                abstract.find(qn('w:lvl')).find(qn('w:' + tag)).set(qn('w:val'), value)
                with self.assertRaisesRegex(ValueError, 'Onbevestigde'):
                    self.parse(document)

    def test_three_paragraphs_in_same_instance_continue(self):
        records = self.parse(self.regular_document([(801, 1)] * 3))
        self.assertEqual([r.code for r in records], [f'I-Test-a LPD {n}' for n in range(1, 4)])

    def test_nonzero_regular_level_is_rejected(self):
        document = self.regular_document([(801, 1)])
        document.paragraphs[0]._p.pPr.find(qn('w:numPr')).find(qn('w:ilvl')).set(qn('w:val'), '1')
        with self.assertRaisesRegex(ValueError, 'Onbevestigde'):
            self.parse(document)

    def test_abstract_start_five_without_override_is_rejected(self):
        with self.assertRaisesRegex(ValueError, 'startwaarde'):
            self.parse(self.regular_document([(801, None)], base_start='5'))

    def test_abstract_start_one_without_override_continues(self):
        records = self.parse(self.regular_document([(801, None)] * 3))
        self.assertEqual([r.code for r in records], [f'I-Test-a LPD {n}' for n in range(1, 4)])

    def test_override_one_takes_precedence_over_abstract_five(self):
        records = self.parse(self.regular_document([(801, 1)] * 2, base_start='5'))
        self.assertEqual([r.code for r in records], ['I-Test-a LPD 1', 'I-Test-a LPD 2'])

    def test_omitted_abstract_start_defaults_to_unsupported_zero(self):
        with self.assertRaisesRegex(ValueError, 'effectieve startwaarde'):
            self.parse(self.regular_document([(801, None)], base_start=None))
        self.assertEqual(len(self.parse(self.regular_document([(801, 1)], base_start=None))), 1)

    def test_invalid_abstract_start_metadata_is_rejected(self):
        for value in ['', 'abc', '1.5']:
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, 'startmetadata'):
                self.parse(self.regular_document([(801, None)], base_start=value))
        document = self.regular_document([(801, None)])
        abstract = next(a for a in document.part.numbering_part.element.findall(qn('w:abstractNum')) if a.get(qn('w:abstractNumId')) == '800')
        del abstract.find(qn('w:lvl')).find(qn('w:start')).attrib[qn('w:val')]
        with self.assertRaisesRegex(ValueError, 'startmetadata'):
            self.parse(document)

    def test_invalid_override_cannot_fall_back_to_abstract_start(self):
        for value in ['', 'abc', '1.5']:
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, 'startmetadata'):
                self.parse(self.regular_document([(801, value)]))
        document = self.regular_document([(801, 1)])
        num = next(n for n in document.part.numbering_part.element.findall(qn('w:num')) if n.get(qn('w:numId')) == '801')
        del num.find(qn('w:lvlOverride')).find(qn('w:startOverride')).attrib[qn('w:val')]
        with self.assertRaisesRegex(ValueError, 'startmetadata'):
            self.parse(document)

    def test_zero_and_negative_starts_are_outside_supported_sequence(self):
        for value in ['0', '-1']:
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, 'startwaarde'):
                self.parse(self.regular_document([(801, None)], base_start=value))

    def test_level_override_without_start_uses_abstract_start(self):
        document = self.regular_document([(801, 1)], base_start='5')
        num = next(n for n in document.part.numbering_part.element.findall(qn('w:num')) if n.get(qn('w:numId')) == '801')
        override = num.find(qn('w:lvlOverride'))
        override.remove(override.find(qn('w:startOverride')))
        with self.assertRaisesRegex(ValueError, 'effectieve startwaarde'):
            self.parse(document)

    def test_duplicate_start_definitions_are_rejected(self):
        from copy import deepcopy
        for location in ['abstract', 'override']:
            with self.subTest(location=location):
                document = self.regular_document([(801, None if location == 'abstract' else 1)])
                numbering = document.part.numbering_part.element
                if location == 'abstract':
                    abstract = next(a for a in numbering.findall(qn('w:abstractNum')) if a.get(qn('w:abstractNumId')) == '800')
                    parent = abstract.find(qn('w:lvl')); child = parent.find(qn('w:start'))
                else:
                    num = next(n for n in numbering.findall(qn('w:num')) if n.get(qn('w:numId')) == '801')
                    parent = num.find(qn('w:lvlOverride')); child = parent.find(qn('w:startOverride'))
                parent.append(deepcopy(child))
                with self.assertRaisesRegex(ValueError, 'startmetadata'):
                    self.parse(document)
